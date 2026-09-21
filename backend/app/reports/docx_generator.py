import io
from datetime import datetime
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def style_word_table(table, col_widths=None, header_bg="F1F5F9"):
    """
    Applies authentic Word table grid styling with shaded header and explicit column widths.
    Ensures that Word displays clear visible borders and cell padding while keeping cells editable.
    """
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row Shading
    if len(table.rows) > 0:
        for idx, cell in enumerate(table.rows[0].cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_bg}"/>'))

    # Set column widths & cell padding across all rows
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for idx, cell in enumerate(row.cells):
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'<w:top w:w="90" w:type="dxa"/>'
                f'<w:bottom w:w="90" w:type="dxa"/>'
                f'<w:left w:w="120" w:type="dxa"/>'
                f'<w:right w:w="120" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)


class DOCXReportGenerator:
    """Generates an editable Microsoft Word (.docx) Legal Metrology Compliance Inspection Report."""

    def generate(self, scan_data: Dict[str, Any]) -> bytes:
        doc = Document()

        # Set standard margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Header Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_gov = title_p.add_run("GOVERNMENT OF INDIA • MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION\n")
        run_gov.font.size = Pt(11)
        run_gov.font.bold = True
        run_gov.font.color.rgb = RGBColor(15, 23, 42)

        run_dept = title_p.add_run("DEPARTMENT OF LEGAL METROLOGY — INSPECTION & ENFORCEMENT DIVISION\n")
        run_dept.font.size = Pt(10)
        run_dept.font.color.rgb = RGBColor(71, 85, 105)

        run_rep = title_p.add_run("Packaged Commodity Compliance Inspection Report\n")
        run_rep.font.size = Pt(14)
        run_rep.font.bold = True
        run_rep.font.color.rgb = RGBColor(30, 58, 138)

        run_sub = title_p.add_run("Issued under Section 18 & 52 of Legal Metrology Act, 2009 read with Legal Metrology (Packaged Commodities) Rules, 2011")
        run_sub.font.size = Pt(8.5)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph() # Spacing

        # Metadata Table
        meta_table = doc.add_table(rows=3, cols=2)
        panels_count = len(scan_data.get('image_urls') or [1])
        ref_scale = f"{scan_data.get('reference_scale_mm')} mm" if scan_data.get('reference_scale_mm') else 'None Provided'
        created_str = scan_data.get('created_at').strftime('%d-%b-%Y %H:%M UTC') if hasattr(scan_data.get('created_at'), 'strftime') else str(scan_data.get('created_at', 'N/A'))
        meta_data = [
            [f"Inspection ID: {scan_data.get('id', 'N/A')}", f"Date & Time: {created_str}"],
            [f"Product / Brand: {scan_data.get('product_name', 'Commodity')}", f"Category: {scan_data.get('category', 'Packaged Goods')}"],
            [f"Inspecting Officer: {scan_data.get('inspector_name', 'Senior Inspector')}", f"Panels Inspected: {panels_count} Panel(s) | Scale: {ref_scale}"]
        ]
        for r_idx, row in enumerate(meta_table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.text = meta_data[r_idx][c_idx]
                cell.paragraphs[0].runs[0].font.size = Pt(9)
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
        style_word_table(meta_table, [Inches(3.33), Inches(3.34)], header_bg="F8FAFC")

        doc.add_paragraph()

        # Cross-reference directive callout
        cross_refs = scan_data.get("cross_referenced_fields")
        if cross_refs:
            ref_p = doc.add_paragraph()
            ref_run = ref_p.add_run(f"CROSS-PANEL DIRECTIVES: Declarations for {', '.join(cross_refs)} contain cross-panel pointers (e.g. 'SEE BASE / SIDE PANEL') directing officers to secondary packaging surfaces.")
            ref_run.font.italic = True
            ref_run.font.size = Pt(8.5)
            ref_run.font.color.rgb = RGBColor(30, 64, 175)

        # Verdict Paragraph
        verdict = str(scan_data.get("overall_compliance_verdict", "PENDING")).replace("_", " ")
        score = scan_data.get("compliance_score", 0.0)
        v_p = doc.add_paragraph()
        v_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        v_run = v_p.add_run(f"OVERALL VERDICT: {verdict.upper()}  |  COMPLIANCE SCORE: {score}%\n")
        v_run.font.bold = True
        v_run.font.size = Pt(12)
        if "COMPLIANT" in verdict and "NON" not in verdict:
            v_run.font.color.rgb = RGBColor(21, 128, 61)
        elif "NON" in verdict:
            v_run.font.color.rgb = RGBColor(185, 28, 28)
        else:
            v_run.font.color.rgb = RGBColor(180, 83, 9)

        # ----------------------------------------------------------------------
        # Section 1: Rule-by-rule Checklist (Word Table)
        # ----------------------------------------------------------------------
        doc.add_heading("1. Rule-by-Rule Compliance Checklist (Rules 2011)", level=2)
        rule_results = scan_data.get("rule_results", [])
        if rule_results:
            rule_table = doc.add_table(rows=1, cols=4)
            headers = ["Rule & Legal Citation", "Verdict", "Declared Text Found", "Inspector Findings"]
            hdr_cells = rule_table.rows[0].cells
            for idx, text in enumerate(headers):
                hdr_cells[idx].text = text
                hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                hdr_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

            for rule in rule_results:
                row_cells = rule_table.add_row().cells
                panel_str = f" [Panel {rule.get('image_index', 0) + 1}]" if rule.get('image_index', 0) > 0 else ""
                row_cells[0].text = f"{rule.get('title')}{panel_str}\n({rule.get('clause_reference')})"
                row_cells[1].text = str(rule.get("status", "N/A")).upper()
                row_cells[2].text = str(rule.get("extracted_value") or "Not Found")
                row_cells[3].text = str(rule.get("remarks", ""))

                # Status coloring
                stat = rule.get("status", "")
                if stat == "PASS":
                    row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
                elif stat == "FAIL":
                    row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
                else:
                    row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 83, 9)
                row_cells[1].paragraphs[0].runs[0].font.bold = True

                for cell in row_cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8.5)

            style_word_table(rule_table, [Inches(2.0), Inches(0.85), Inches(1.8), Inches(2.02)])

        doc.add_paragraph()

        # ----------------------------------------------------------------------
        # Section 2: Mandatory Declarations (Table 1.0)
        # ----------------------------------------------------------------------
        doc.add_heading("2. Extracted Mandatory Declarations (Table 1.0 - Auditable Evidence)", level=2)
        extracted = scan_data.get("extracted_data", {})
        master = scan_data.get("master_report", {}) or extracted
        master_mand = master.get("mandatory_declarations", {})
        master_mfg = master.get("manufacturer_details", {})

        dec_table = doc.add_table(rows=1, cols=5)
        dec_headers = ["Declaration Field", "Declared Value Found", "Status", "Confidence", "Statutory Audit Remarks"]
        for idx, text in enumerate(dec_headers):
            dec_table.rows[0].cells[idx].text = text
            dec_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
            dec_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
            dec_table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

        field_defs = [
            ("product_name", "Product / Brand Name", "found", "Identified from principal display panel."),
            ("generic_name", "Common / Generic Commodity Name (Rule 6(1)(b))", "missing", "Rule 6(1)(b) verification."),
            ("manufacturer_name", "Manufacturer / Packer Legal Name (Rule 6(1)(a))", "invalid", "Rule 6(1)(a) verification."),
            ("manufacturer_address", "Complete Physical Factory Address (Rule 6(1)(a))", "invalid", "Rule 6(1)(a) factory premises verification."),
            ("net_quantity", "Net Quantity & Standard SI Units (Rule 6(1)(c) & Rule 13)", "invalid", "Rule 6(1)(c) & Rule 13 audit."),
            ("mrp", "Maximum Retail Price [MRP] (Rule 6(1)(e))", "invalid", "Rule 6(1)(e) retail price declaration."),
            ("unit_sale_price", "Unit Sale Price [USP] (Rule 6(1)(e) Amendment)", "missing", "Rule 6(1)(e) USP mandate."),
            ("mfg_date", "Month & Year of Manufacture / Packing (Rule 6(1)(d))", "missing", "Rule 6(1)(d) date declaration."),
            ("consumer_care", "Consumer Care Contact Details (Rule 6(1)(f))", "missing", "Rule 6(1)(f) grievance contact."),
            ("country_of_origin", "Country of Origin Declaration (Rule 6(1)(g))", "missing", "Rule 6(1)(g) origin declaration.")
        ]

        edited_fields = scan_data.get("edited_fields") or {}

        for key, label, def_stat, def_rem in field_defs:
            f = extracted.get(key) if isinstance(extracted.get(key), dict) else None
            is_edited = (key in edited_fields) or (f and f.get("is_manually_edited"))
            val = (edited_fields.get(key, {}).get("edited_value")) or (f.get("value") if f else None) or master_mand.get(key, {}).get("declared_value") or (scan_data.get("product_name") if key == "product_name" else None)
            if not val and key == "manufacturer_name":
                val = master_mfg.get("resolved_manufacturer_name")
            if not val and key == "manufacturer_address":
                val = master_mfg.get("resolved_address")

            stat = (f.get("status") if f else None) or master_mand.get(key, {}).get("compliance_status") or (master_mfg.get("compliance_status") if "manufacturer" in key else None) or ("found" if val else def_stat)
            rem = (f.get("validation_remarks") if f else None) or master_mand.get(key, {}).get("compliance_remarks") or (master_mfg.get("compliance_remarks") if "manufacturer" in key else None) or def_rem
            conf = (f.get("confidence") if f else None) or 0.95

            is_pass = stat in ["valid", "found", "COMPLIANT"]
            is_rev = stat in ["review_required", "NEEDS_REVIEW", "NEEDS_MANUAL_INSPECTION"]
            disp_stat = "PASS" if is_pass else ("REVIEW" if is_rev else "FAIL")

            row_cells = dec_table.add_row().cells
            p_label = row_cells[0].paragraphs[0]
            p_label.text = label
            if is_edited:
                run_tag = p_label.add_run("  [MANUAL EDIT]")
                run_tag.font.size = Pt(7.5)
                run_tag.font.bold = True
                run_tag.font.color.rgb = RGBColor(180, 83, 9)

            row_cells[1].text = str(val or "Not Declared on Pack")
            row_cells[2].text = disp_stat
            row_cells[3].text = f"{int(round(conf * 100))}%"
            row_cells[4].text = str(rem or "")

            if is_pass:
                row_cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
            elif is_rev:
                row_cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 83, 9)
            else:
                row_cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
            row_cells[2].paragraphs[0].runs[0].font.bold = True

            for cell in row_cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run != getattr(p_label, 'runs', [None])[-1] or not is_edited:
                            run.font.size = Pt(8.5)

        style_word_table(dec_table, [Inches(1.8), Inches(1.7), Inches(0.75), Inches(0.7), Inches(1.72)])

        doc.add_paragraph()

        # ----------------------------------------------------------------------
        # Section 3: Manufacturer, Packer & Importer Details Table
        # ----------------------------------------------------------------------
        doc.add_heading("3. Manufacturer, Packer & Importer Traceability Details (Rule 6(1)(a) & Rule 10)", level=2)
        mfg_table = doc.add_table(rows=1, cols=4)
        mfg_headers = ["Traceability Parameter", "Resolved Declaration Details", "Verification Source", "Compliance Findings"]
        for idx, text in enumerate(mfg_headers):
            mfg_table.rows[0].cells[idx].text = text
            mfg_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
            mfg_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
            mfg_table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

        mfg_rows_data = [
            ("Manufacturer / Packer Name", master_mfg.get("resolved_manufacturer_name") or "Not Declared", master_mfg.get("resolution_method", "Packaging Text"), "Physical factory identity recorded."),
            ("Factory Premises Address", master_mfg.get("resolved_address") or "Not Declared", master_mfg.get("resolution_method", "Packaging Text"), master_mfg.get("compliance_remarks", "Factory address verification under Rule 6(1)(a).")),
            ("QR Portal / Digital Link", master.get("qr_code_detected_url") or master_mfg.get("target_url") or "None on Packaging", "Optical QR Sweep", f"Batch code used: '{master_mfg.get('batch_code_used', 'N/A')}'" if master_mfg.get('batch_code_used') else "Direct label declaration."),
            ("Country of Origin", master_mand.get("country_of_origin", {}).get("declared_value") or "India", "Rule 6(1)(g)", master_mand.get("country_of_origin", {}).get("compliance_remarks", "Country of origin verified."))
        ]

        for p_label, p_val, p_src, p_find in mfg_rows_data:
            r_cells = mfg_table.add_row().cells
            r_cells[0].text = p_label
            r_cells[1].text = str(p_val)
            r_cells[2].text = str(p_src)
            r_cells[3].text = str(p_find)
            for c in r_cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)

        style_word_table(mfg_table, [Inches(1.7), Inches(2.2), Inches(1.1), Inches(1.67)])

        doc.add_paragraph()

        # ----------------------------------------------------------------------
        # Section 4: USP Cross-Verification (Table 2.0)
        # ----------------------------------------------------------------------
        usp = master.get("usp_cross_verification") or extracted.get("usp_cross_verification")
        if usp and (usp.get("mrp_numeric") or usp.get("calculated_usp_per_base_unit")):
            doc.add_heading("4. Unit Sale Price (USP) Mathematical Cross-Verification (Rule 6(1)(e))", level=2)
            usp_table = doc.add_table(rows=1, cols=5)
            u_hdrs = ["Declared MRP", "Net Quantity", "Declared USP", "Calculated Base USP", "Verdict"]
            for idx, text in enumerate(u_hdrs):
                usp_table.rows[0].cells[idx].text = text
                usp_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
                usp_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                usp_table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

            u_row = usp_table.add_row().cells
            u_row[0].text = f"Rs. {usp.get('mrp_numeric', 'N/A')}"
            u_row[1].text = f"{usp.get('quantity_numeric', 'N/A')} {usp.get('quantity_unit', '')}"
            u_row[2].text = str(usp.get('declared_usp_raw') or (f"Rs. {usp.get('declared_usp_numeric')}" if usp.get('declared_usp_numeric') else 'Not declared'))
            u_row[3].text = f"Rs. {usp.get('calculated_usp_per_base_unit', 'N/A')} {usp.get('calculated_usp_unit', '')}"
            u_row[4].text = "COMPLIANT" if usp.get('is_compliant') else "VIOLATION"

            if usp.get('is_compliant'):
                u_row[4].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
            else:
                u_row[4].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
            u_row[4].paragraphs[0].runs[0].font.bold = True

            for c in u_row:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)

            style_word_table(usp_table, [Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.6), Inches(1.17)])

            rem_p = doc.add_paragraph(f"Statutory Analysis: {usp.get('remarks', '')}")
            rem_p.runs[0].font.size = Pt(8.5)
            rem_p.runs[0].font.italic = True
            if usp.get("violations"):
                inf_p = doc.add_paragraph(f"Statutory Infractions: {'; '.join(usp.get('violations'))}")
                inf_p.runs[0].font.size = Pt(8.5)
                inf_p.runs[0].font.bold = True
                inf_p.runs[0].font.color.rgb = RGBColor(185, 28, 28)

        # ----------------------------------------------------------------------
        # Section 5: Second Schedule Shrinkflation Audit (Table 3.0)
        # ----------------------------------------------------------------------
        shrink = master.get("second_schedule_shrinkflation_audit") or extracted.get("second_schedule_shrinkflation_audit")
        if shrink and shrink.get("is_second_schedule_commodity"):
            doc.add_heading("5. Second Schedule Prescribed Pack Sizes & Anti-Shrinkflation Audit (Rule 5)", level=2)
            s_table = doc.add_table(rows=1, cols=5)
            s_hdrs = ["Prescribed Commodity", "Declared Pack Size", "Nearest Standard", "Deviation / Downsizing", "Verdict"]
            for idx, text in enumerate(s_hdrs):
                s_table.rows[0].cells[idx].text = text
                s_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
                s_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                s_table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

            s_row = s_table.add_row().cells
            s_row[0].text = str(shrink.get('matched_commodity_category', 'Not Restricted'))
            s_row[1].text = str(shrink.get('declared_quantity', 'N/A'))
            s_row[2].text = str(shrink.get('nearest_standard_pack_size') or 'N/A')
            s_row[3].text = f"{shrink.get('shrinkage_percentage', 0)}%"
            s_row[4].text = "PASS" if shrink.get('is_compliant') else "FAIL"

            if shrink.get('is_compliant'):
                s_row[4].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
            else:
                s_row[4].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
            s_row[4].paragraphs[0].runs[0].font.bold = True

            for c in s_row:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)

            style_word_table(s_table, [Inches(1.8), Inches(1.2), Inches(1.2), Inches(1.3), Inches(1.17)])

            s_rem = doc.add_paragraph(f"Findings: {shrink.get('remarks', '')}")
            s_rem.runs[0].font.size = Pt(8.5)
            s_rem.runs[0].font.italic = True

        # ----------------------------------------------------------------------
        # Section 6: Physical & Visual Metrology Audit (Rule 7 & 9 - Table 4.0)
        # ----------------------------------------------------------------------
        metro = master.get("visual_and_metrology_audit") or extracted.get("visual_and_metrology_audit")
        if metro and (metro.get("rule_7_font_size") or metro.get("rule_9_color_contrast")):
            doc.add_heading("6. Physical & Visual Metrology Audit (Rule 7 & 9)", level=2)
            m_table = doc.add_table(rows=1, cols=5)
            m_hdrs = ["Metrology Parameter", "Measured Value", "Statutory Benchmark", "Verdict", "Remarks"]
            for idx, text in enumerate(m_hdrs):
                m_table.rows[0].cells[idx].text = text
                m_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
                m_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                m_table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

            if metro.get("rule_7_font_size"):
                f_data = metro["rule_7_font_size"]
                f_row = m_table.add_row().cells
                f_label_p = f_row[0].paragraphs[0]
                f_label_p.text = "Table-I Numeral & Letter Height (Rule 7(2))"
                if ("font_size" in edited_fields) or bool(extracted.get("font_size", {}).get("is_manually_edited")):
                    r_tag = f_label_p.add_run("  [MANUAL EDIT]")
                    r_tag.font.size = Pt(7.5)
                    r_tag.font.bold = True
                    r_tag.font.color.rgb = RGBColor(180, 83, 9)
                f_row[1].text = f"{f_data.get('measured_numeral_height_mm')} mm" if f_data.get('measured_numeral_height_mm') else "Optical Inspection"
                f_row[2].text = f"Min {f_data.get('statutory_min_height_mm', 1.0)} mm"
                f_row[3].text = "PASS" if f_data.get('is_compliant') else "FAIL"
                f_row[4].text = str(f_data.get('remarks') or "")
                if f_data.get('is_compliant'):
                    f_row[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
                else:
                    f_row[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
                f_row[3].paragraphs[0].runs[0].font.bold = True

            if metro.get("rule_9_color_contrast"):
                c_data = metro["rule_9_color_contrast"]
                c_row = m_table.add_row().cells
                c_label_p = c_row[0].paragraphs[0]
                c_label_p.text = "Background Color Contrast (Rule 9(1)(b))"
                if ("color_contrast" in edited_fields) or bool(extracted.get("color_contrast", {}).get("is_manually_edited")):
                    r_tag = c_label_p.add_run("  [MANUAL EDIT]")
                    r_tag.font.size = Pt(7.5)
                    r_tag.font.bold = True
                    r_tag.font.color.rgb = RGBColor(180, 83, 9)
                c_row[1].text = f"Ratio: {c_data.get('contrast_ratio', 'N/A')}"
                c_row[2].text = "Min 3.0:1 Ratio"
                c_row[3].text = "PASS" if c_data.get('is_compliant') else "FAIL"
                c_row[4].text = str(c_data.get('remarks') or "")
                if c_data.get('is_compliant'):
                    c_row[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)
                else:
                    c_row[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
                c_row[3].paragraphs[0].runs[0].font.bold = True

            for row in m_table.rows[1:]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(8.5)

            style_word_table(m_table, [Inches(2.0), Inches(1.2), Inches(1.2), Inches(0.9), Inches(1.37)])

        # ----------------------------------------------------------------------
        # Section 7: Statutory Violations & Enforcement Penalties
        # ----------------------------------------------------------------------
        summary = master.get("statutory_summary") or extracted.get("statutory_summary") or []
        actual_violations = [
            item for item in summary
            if "comply strictly" not in item.lower() and "no statutory" not in item.lower() and "nil" not in item.lower()
        ]
        if actual_violations:
            doc.add_heading("7. Statutory Violations & Enforcement Penalties (Legal Metrology Act, 2009)", level=2)
            for idx, item in enumerate(actual_violations):
                p = doc.add_paragraph(f"Violation {idx + 1}: {item}", style='List Bullet')
                p.runs[0].font.size = Pt(8.5)
                p.runs[0].font.color.rgb = RGBColor(185, 28, 28)
        else:
            doc.add_heading("7. Statutory Enforcement Summary (Legal Metrology Act, 2009)", level=2)
            p = doc.add_paragraph("Nil Infractions Recorded: All mandatory packaging declarations comply strictly with the Legal Metrology (Packaged Commodities) Rules, 2011. No statutory penalties under Section 36(1) or Section 39 applicable.")
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(21, 128, 61)
            p.runs[0].font.bold = True

        # ----------------------------------------------------------------------
        # Section 8: Inspector Notes
        # ----------------------------------------------------------------------
        notes = scan_data.get("inspector_notes")
        if notes:
            doc.add_heading("8. Inspecting Officer Observations & Evidence Notes", level=2)
            p = doc.add_paragraph(notes)
            p.runs[0].font.size = Pt(9)

        # Statutory Explanatory Notice for Badges
        p_notice = doc.add_paragraph()
        p_notice.paragraph_format.space_before = Pt(12)
        p_notice.paragraph_format.space_after = Pt(12)
        run_notice_title = p_notice.add_run("STATUTORY NOTICE ON EVIDENTIARY BADGES: ")
        run_notice_title.bold = True
        run_notice_title.font.size = Pt(8)
        run_notice_title.font.color.rgb = RGBColor(180, 83, 9)
        run_notice_body = p_notice.add_run(
            "Declarations marked with the amber [MANUAL EDIT] tag represent statutory parameters physically inspected, "
            "verified, and updated by the authorized enforcement officer under Section 15 of the Legal Metrology Act, 2009, "
            "superseding automated OCR extractions. Entries without this badge represent automated multimodal detections."
        )
        run_notice_body.font.size = Pt(8)
        run_notice_body.font.italic = True
        run_notice_body.font.color.rgb = RGBColor(100, 116, 139)

        # Signature Block
        doc.add_paragraph("\n")
        sig_p = doc.add_paragraph(f"{scan_data.get('inspector_name', 'Authorized Legal Metrology Inspector')}                 Official Seal & Digital Stamp\nLegal Metrology Department                                  Ministry of Consumer Affairs")
        sig_p.runs[0].font.size = Pt(8.5)

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

docx_report_generator = DOCXReportGenerator()
